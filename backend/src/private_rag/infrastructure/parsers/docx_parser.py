"""DOCX parser built on python-docx."""

from __future__ import annotations

from pathlib import Path

import docx

from private_rag.application.dto.parsing import ParsedDocument, TextBlock
from private_rag.application.exceptions import (
    DocumentExtractionError,
    NoExtractableTextError,
)
from private_rag.application.services.text_normalization import normalize_text


class DocxParser:
    """Extracts paragraph text from a Word (.docx) document.

    DOCX has no reliable page model, so all text is returned as a single block
    without page numbers.
    """

    def parse(self, file_path: Path) -> ParsedDocument:
        try:
            document = docx.Document(str(file_path))
        except Exception as exc:  # python-docx raises a variety of errors
            raise DocumentExtractionError(f"Could not open DOCX: {exc}") from exc

        paragraphs = [para.text for para in document.paragraphs if para.text.strip()]
        normalized = normalize_text("\n\n".join(paragraphs))
        if not normalized:
            raise NoExtractableTextError("DOCX contains no readable text")

        return ParsedDocument(blocks=[TextBlock(text=normalized, page_number=None)])
