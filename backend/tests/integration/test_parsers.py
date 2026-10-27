"""Integration tests for document parsers."""

from __future__ import annotations

from pathlib import Path

import docx
import pymupdf
import pytest

from private_rag.application.exceptions import (
    DocumentExtractionError,
    NoExtractableTextError,
    UnsupportedParserError,
)
from private_rag.infrastructure.parsers.pdf_parser import PdfParser
from private_rag.infrastructure.parsers.registry import ParserRegistry
from private_rag.infrastructure.parsers.text_parser import PlainTextParser


class TestPlainTextParser:
    def test_parses_and_normalizes(self, tmp_path: Path) -> None:
        path = tmp_path / "a.txt"
        path.write_text("Hello   world\r\n\r\n\r\n\r\nSecond   paragraph", encoding="utf-8")

        result = PlainTextParser().parse(path)

        assert result.text == "Hello world\n\nSecond paragraph"
        assert result.blocks[0].page_number is None

    def test_empty_raises(self, tmp_path: Path) -> None:
        path = tmp_path / "a.txt"
        path.write_text("   \n\n  ", encoding="utf-8")
        with pytest.raises(NoExtractableTextError):
            PlainTextParser().parse(path)

    def test_invalid_utf8_raises(self, tmp_path: Path) -> None:
        path = tmp_path / "a.txt"
        path.write_bytes(b"\xff\xfe\x00\x01bad")
        with pytest.raises(DocumentExtractionError):
            PlainTextParser().parse(path)


class TestDocxParser:
    def test_parses_paragraphs(self, tmp_path: Path) -> None:
        path = tmp_path / "a.docx"
        document = docx.Document()
        document.add_paragraph("First paragraph.")
        document.add_paragraph("Second paragraph.")
        document.save(str(path))

        result = ParserRegistry().parse(path, ".docx")

        assert "First paragraph." in result.text
        assert "Second paragraph." in result.text


class TestPdfParser:
    def test_parses_pages_with_numbers(self, tmp_path: Path) -> None:
        path = tmp_path / "a.pdf"
        document = pymupdf.open()
        page1 = document.new_page()
        page1.insert_text((72, 72), "Alpha content on page one")
        page2 = document.new_page()
        page2.insert_text((72, 72), "Beta content on page two")
        document.save(str(path))
        document.close()

        result = PdfParser().parse(path)

        assert len(result.blocks) == 2
        assert result.blocks[0].page_number == 1
        assert result.blocks[1].page_number == 2
        assert "Alpha" in result.blocks[0].text

    def test_scanned_pdf_without_text_raises(self, tmp_path: Path) -> None:
        path = tmp_path / "scanned.pdf"
        document = pymupdf.open()
        document.new_page()  # empty page, no text layer
        document.save(str(path))
        document.close()

        with pytest.raises(NoExtractableTextError):
            PdfParser().parse(path)

    def test_corrupted_pdf_raises(self, tmp_path: Path) -> None:
        path = tmp_path / "broken.pdf"
        path.write_bytes(b"%PDF-1.7 totally broken")
        with pytest.raises(DocumentExtractionError):
            PdfParser().parse(path)


def test_registry_rejects_unsupported_extension(tmp_path: Path) -> None:
    with pytest.raises(UnsupportedParserError):
        ParserRegistry().parse(tmp_path / "x.rtf", ".rtf")
